"""Messy-file fixtures for the guided upload layer.

Every case builds ONE file (csv or xlsx, written programmatically) that is a
realistic way people hand over sales history, together with what the guide must
say about it and what the file must become once the questions are answered:

* ``findings``  - codes that must appear in the first report;
* ``asks``      - question codes, in order, the walk-through must meet;
* ``answers``   - {question code: option id} the simulated user gives;
* ``fixes``     - fix codes the final report must contain;
* ``clean``     - the canonical (sku, date, demand) rows the file must produce
                  after the fixes, i.e. what its clean equivalent contains.

Shared by ForecastingCore/tests/test_guidance.py and the backend endpoint
tests (loaded by path), so the two layers are judged on the same files.
"""

from __future__ import annotations

import csv
import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Optional

SKUS = ["SKU-001", "SKU-002", "SKU-003", "SKU-004"]
START = dt.date(2025, 1, 6)                      # a Monday
WEEKS = 26


def demand(i: int, k: int) -> int:
    """Deterministic noisy weekly demand: not monotone, not constant."""
    return 20 + ((k * 7 + i * 5 + (k * k) % 11) % 17)


def base(skus=SKUS, weeks=WEEKS, start=START, step=7) -> list[tuple]:
    out = []
    for i, s in enumerate(skus):
        for k in range(weeks):
            out.append((s, start + dt.timedelta(days=step * k), float(demand(i, k))))
    return out


def month_start(k: int, year: int = 2024) -> dt.date:
    return dt.date(year + (k // 12), k % 12 + 1, 1)


def clean_of(rows) -> list[tuple]:
    return [(s, d.isoformat() if hasattr(d, "isoformat") else d, float(v)) for s, d, v in rows]


# ── writers ──────────────────────────────────────────────────────────────────

def write_csv(path: Path, rows: list[list], sep: str = ",", encoding: str = "utf-8") -> Path:
    with open(path, "w", newline="", encoding=encoding) as fh:
        w = csv.writer(fh, delimiter=sep)
        for r in rows:
            w.writerow(["" if c is None else c for c in r])
    return path


def write_xlsx(path: Path, sheets: dict[str, list[list]], merges: Optional[dict] = None) -> Path:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(list(r))
        for rng in (merges or {}).get(name, []):
            ws.merge_cells(rng)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, dt.datetime):
                    c.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(c.value, dt.date):
                    c.number_format = "yyyy-mm-dd"
    wb.save(path)
    return path


def fmt_dmy(d: dt.date) -> str:
    return f"{d.day:02d}/{d.month:02d}/{d.year}"


ES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
EN = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


@dataclass
class Case:
    id: str
    build: Callable[[Path], Path]
    clean: Optional[list] = None
    verdict: str = "ready"
    findings: list = field(default_factory=list)
    asks: list = field(default_factory=list)
    answers: dict = field(default_factory=dict)
    fixes: list = field(default_factory=list)
    mapping: dict = field(default_factory=dict)
    absent: list = field(default_factory=list)       # finding codes that must NOT appear
    zero_fixes: bool = False
    extra: dict = field(default_factory=dict)


CASES: list[Case] = []


def case(**kw):
    def deco(fn):
        CASES.append(Case(id=fn.__name__, build=fn, **kw))
        return fn
    return deco


B = base()
STD = {"sku": "sku", "date": "fecha", "demand": "cantidad"}


def _std_rows(rows=B, header=("sku", "fecha", "cantidad"), dfmt=lambda d: d.isoformat(),
              vfmt=lambda v: int(v)):
    return [list(header)] + [[s, dfmt(d), vfmt(v)] for s, d, v in rows]


# ── clean files: ZERO questions, ZERO fixes ──────────────────────────────────

@case(clean=clean_of(B), zero_fixes=True, mapping=STD)
def clean_csv(tmp: Path):
    return write_csv(tmp / "clean_csv.csv", _std_rows())


@case(clean=clean_of(B), zero_fixes=True, mapping=STD)
def clean_xlsx_native_dates(tmp: Path):
    rows = [["sku", "fecha", "cantidad"]] + [[s, d, int(v)] for s, d, v in B]
    return write_xlsx(tmp / "clean_xlsx.xlsx", {"Hoja1": rows})


@case(clean=clean_of(B), zero_fixes=True,
      mapping={"sku": "Codigo", "date": "Fecha", "demand": "Unidades"})
def clean_semicolon_dayfirst(tmp: Path):
    rows = [["Codigo", "Fecha", "Unidades"]] + [[s, fmt_dmy(d), int(v)] for s, d, v in B]
    return write_csv(tmp / "clean_semicolon.csv", rows, sep=";")


@case(clean=clean_of(B), zero_fixes=True, mapping=STD)
def clean_with_name_and_extra_columns(tmp: Path):
    rows = [["sku", "nombre", "fecha", "cantidad", "precio", "stock"]]
    for s, d, v in B:
        rows.append([s, f"Producto {s[-1]}", d.isoformat(), int(v), 12.5, 40])
    return write_csv(tmp / "clean_extra.csv", rows)


@case(clean=clean_of(B), zero_fixes=True,
      mapping={"sku": "producto", "date": "fecha", "demand": "cantidad"})
def clean_quantity_price_total(tmp: Path):
    rows = [["fecha", "producto", "cantidad", "precio", "total"]]
    for s, d, v in B:
        rows.append([d.isoformat(), s, int(v), 5, int(v) * 5])
    return write_csv(tmp / "clean_total.csv", rows)


@case(clean=clean_of(B), zero_fixes=True,
      mapping={"sku": "sku", "date": "fecha", "demand": "ventas"})
def clean_ventas_only(tmp: Path):
    return write_csv(tmp / "clean_ventas.csv", _std_rows(header=("sku", "fecha", "ventas")))


_DAILY = [(s, dt.date(2025, 3, 1) + dt.timedelta(days=k), float(demand(i, k)))
          for i, s in enumerate(SKUS[:3]) for k in range(60)]


_DAILY2 = [(s, d, v * mult) for s, d, v in _DAILY for mult in (1, 2)]


@case(clean=clean_of(_DAILY2), zero_fixes=True,
      mapping={"sku": "sku", "date": "fecha", "demand": "unidades", "store": "tienda"},
      extra={"granularity": "daily"})
def clean_daily_two_stores(tmp: Path):
    rows = [["sku", "tienda", "fecha", "unidades"]]
    for s, d, v in _DAILY:
        for store, mult in (("Norte", 1), ("Sur", 2)):
            rows.append([s, store, d.isoformat(), int(v) * mult])
    return write_csv(tmp / "clean_daily.csv", rows)


_MONTHLY = [(s, month_start(k), float(demand(i, k) * 10)) for i, s in enumerate(SKUS[:3])
            for k in range(24)]


@case(clean=clean_of(_MONTHLY), zero_fixes=True, mapping=STD, extra={"granularity": "monthly"})
def clean_monthly(tmp: Path):
    return write_csv(tmp / "clean_monthly.csv", _std_rows(_MONTHLY))


_RET = [(s, d, v if (i + k) % 9 else -3.0)
        for k, (s, d, v) in enumerate(B) for i in (0,)]


@case(clean=clean_of(_RET), zero_fixes=True, mapping=STD, findings=["demand_negatives"])
def clean_with_returns(tmp: Path):
    return write_csv(tmp / "returns.csv", _std_rows(_RET))


@case(clean=clean_of(B), zero_fixes=True, mapping=STD, extra={"rowid": True},
      absent=["sku_column_choice"])
def sku_row_id_column_ignored(tmp: Path):
    rows = [["linea", "sku", "fecha", "cantidad"]]
    for n, (s, d, v) in enumerate(B, 1):
        rows.append([n, s, d.isoformat(), int(v)])
    return write_csv(tmp / "rowid.csv", rows)


@case(clean=clean_of(B), zero_fixes=True,
      mapping={"sku": "sku", "date": "cantidad", "demand": "fecha"},
      findings=["columns_swapped"])
def swapped_columns_by_content(tmp: Path):
    rows = [["sku", "cantidad", "fecha"]] + [[s, d.isoformat(), int(v)] for s, d, v in B]
    return write_csv(tmp / "swapped.csv", rows)


_DUP = B + [("SKU-001", START, 999.0), ("SKU-002", START + dt.timedelta(days=7), 5.0)]


@case(clean=clean_of(_DUP), zero_fixes=True, mapping=STD, findings=["conflicting_duplicates"])
def duplicates_conflicting(tmp: Path):
    rows = _std_rows()
    rows += [["SKU-001", START.isoformat(), 999], ["SKU-002", (START + dt.timedelta(days=7)).isoformat(), 5]]
    return write_csv(tmp / "dups.csv", rows)


# ── structure: titles, totals, sheets, wide ──────────────────────────────────

@case(clean=clean_of(B), findings=["header_row_offset"], fixes=["header_row"], mapping=STD)
def title_rows_merged_xlsx(tmp: Path):
    rows = [["Reporte de ventas - Distribuidora Andina", None, None],
            ["Generado el 2025-07-01", None, None], [None, None, None],
            ["sku", "fecha", "cantidad"]] + [[s, d, int(v)] for s, d, v in B]
    return write_xlsx(tmp / "title.xlsx", {"Ventas": rows}, merges={"Ventas": ["A1:C1"]})


@case(clean=clean_of(B), findings=["totals_row"], fixes=["ignore_totals_row"], mapping=STD)
def totals_row_and_note_at_bottom(tmp: Path):
    rows = _std_rows()
    rows.append(["Total", None, int(sum(v for _, _, v in B))])
    rows.append(["Generado por ERP v4"])
    return write_csv(tmp / "totals.csv", rows)


@case(clean=clean_of(B), findings=["totals_row"], fixes=["ignore_totals_row"], mapping=STD)
def subtotal_rows_in_the_middle(tmp: Path):
    rows = [["sku", "fecha", "cantidad"]]
    for s in SKUS:
        part = [(a, d, v) for a, d, v in B if a == s]
        rows += [[a, d.isoformat(), int(v)] for a, d, v in part]
        rows.append(["Subtotal", None, int(sum(v for _, _, v in part))])
    rows.append(["Total general", None, int(sum(v for _, _, v in B))])
    return write_csv(tmp / "subtotals.csv", rows)


@case(clean=clean_of(B), findings=["sheet_auto_selected"], fixes=["use_sheet"], mapping=STD,
      extra={"selected_sheet": "Ventas"})
def multi_sheet_sales_is_second(tmp: Path):
    resumen = [["Resumen"], ["Productos", 4], ["Semanas", 26]]
    ventas = [["sku", "fecha", "cantidad"]] + [[s, d, int(v)] for s, d, v in B]
    catalogo = [["sku", "nombre", "categoria"]] + [[s, f"Prod {s}", "A"] for s in SKUS]
    return write_xlsx(tmp / "multi.xlsx", {"Resumen": resumen, "Ventas": ventas,
                                           "Catalogo": catalogo})


@case(clean=clean_of(B), findings=["sheet_auto_selected"], zero_fixes=True, mapping=STD,
      extra={"selected_sheet": "Ventas"})
def multi_sheet_sales_is_first(tmp: Path):
    ventas = [["sku", "fecha", "cantidad"]] + [[s, d, int(v)] for s, d, v in B]
    catalogo = [["sku", "nombre", "categoria"]] + [[s, f"Prod {s}", "A"] for s in SKUS]
    return write_xlsx(tmp / "multi_first.xlsx", {"Ventas": ventas, "Catalogo": catalogo})


_B24 = base(weeks=26, start=dt.date(2024, 1, 1))
_B25 = base(weeks=26, start=dt.date(2025, 1, 6))


@case(clean=clean_of(_B25), asks=["sheet_choice"], answers={"sheet_choice": "Ventas 2025"},
      fixes=["use_sheet"], mapping=STD)
def multi_sheet_two_equal_candidates(tmp: Path):
    s24 = [["sku", "fecha", "cantidad"]] + [[s, d, int(v)] for s, d, v in _B24]
    s25 = [["sku", "fecha", "cantidad"]] + [[s, d, int(v)] for s, d, v in _B25]
    return write_xlsx(tmp / "twosheets.xlsx", {"Ventas 2024": s24, "Ventas 2025": s25})


def _wide(headers: list, cell=lambda v: int(v), weeks=6, extra_total_col=False, total_row=False):
    skus = SKUS
    rows = [["sku"] + headers + (["Total"] if extra_total_col else [])]
    for i, s in enumerate(skus):
        vals = [float(demand(i, k)) for k in range(len(headers))]
        rows.append([s] + [cell(v) for v in vals] + ([int(sum(vals))] if extra_total_col else []))
    if total_row:
        rows.append(["Total"] + [None] * len(headers) + ([None] if extra_total_col else []))
    return rows


_W_MONTHS = [month_start(k, 2025) for k in range(6)]
_WIDE_CLEAN = [(s, d, float(demand(i, k))) for i, s in enumerate(SKUS)
               for k, d in enumerate(_W_MONTHS)]


@case(clean=clean_of(_WIDE_CLEAN), findings=["wide_format"], fixes=["unpivot"],
      mapping={"sku": "sku", "date": "date", "demand": "demand"})
def wide_datetime_headers_xlsx(tmp: Path):
    rows = _wide([dt.datetime(d.year, d.month, d.day) for d in _W_MONTHS])
    return write_xlsx(tmp / "wide_dt.xlsx", {"Ventas": rows})


@case(clean=clean_of(_WIDE_CLEAN), findings=["wide_format"], fixes=["unpivot"],
      mapping={"sku": "sku", "date": "date", "demand": "demand"})
def wide_spanish_month_year_headers(tmp: Path):
    rows = _wide([f"{ES[d.month - 1]}-25" for d in _W_MONTHS])
    return write_csv(tmp / "wide_es.csv", rows)


@case(clean=clean_of(_WIDE_CLEAN), findings=["wide_year_missing"], asks=["wide_year_missing"],
      answers={"wide_year_missing": "2025"}, fixes=["unpivot"],
      mapping={"sku": "sku", "date": "date", "demand": "demand"})
def wide_month_names_without_year(tmp: Path):
    names = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio"]
    return write_csv(tmp / "wide_months.csv", _wide(names))


@case(clean=clean_of(_WIDE_CLEAN), findings=["wide_format", "totals_row"],
      fixes=["unpivot", "ignore_totals_row", "drop_summary_columns"],
      mapping={"sku": "sku", "date": "date", "demand": "demand"})
def wide_with_total_column_and_row(tmp: Path):
    rows = _wide([d.isoformat() for d in _W_MONTHS], extra_total_col=True, total_row=True)
    return write_csv(tmp / "wide_totals.csv", rows)


# ── dates ────────────────────────────────────────────────────────────────────

_AMB = [(s, month_start(k, 2025), float(demand(i, k))) for i, s in enumerate(SKUS[:3])
        for k in range(12)]


@case(clean=clean_of(_AMB), asks=["date_order_ambiguous"],
      answers={"date_order_ambiguous": "day_first"}, fixes=["read_dates"], mapping=STD)
def date_ambiguous_slash_monthly(tmp: Path):
    return write_csv(tmp / "amb.csv", _std_rows(_AMB, dfmt=fmt_dmy))


_CONF_ROWS = []
for _i, _s in enumerate(SKUS[:2]):
    for _k in range(12):
        _CONF_ROWS.append((_s, dt.date(2025, 1, 1) + dt.timedelta(days=13 * _k), float(demand(_i, _k))))


@case(asks=["date_order_ambiguous", "date_rows_unreadable"],
      answers={"date_order_ambiguous": "day_first", "date_rows_unreadable": "ignore"},
      fixes=["read_dates", "ignore_unreadable_rows"], mapping=STD, extra={"dropped_some": True})
def date_both_orders_in_one_column(tmp: Path):
    rows = [["sku", "fecha", "cantidad"]]
    for n, (s, d, v) in enumerate(_CONF_ROWS):
        rows.append([s, f"{d.day:02d}/{d.month:02d}/{d.year}" if n % 2 == 0
                     else f"{d.month:02d}/{d.day:02d}/{d.year}", int(v)])
    return write_csv(tmp / "conflict.csv", rows)


@case(clean=clean_of(B), findings=["date_text_months"], fixes=["read_dates"], mapping=STD)
def date_spanish_text_months(tmp: Path):
    return write_csv(tmp / "es_dates.csv",
                     _std_rows(dfmt=lambda d: f"{d.day:02d} {ES[d.month - 1]} {d.year}"))


@case(clean=clean_of(B), findings=["date_text_months", "date_two_digit_year"],
      fixes=["read_dates"], mapping=STD)
def date_english_text_two_digit_year(tmp: Path):
    return write_csv(tmp / "en_dates.csv",
                     _std_rows(dfmt=lambda d: f"{d.day:02d}-{EN[d.month - 1]}-{d.year % 100:02d}"))


@case(clean=clean_of(B), findings=["date_excel_serial"], fixes=["read_dates"], mapping=STD)
def date_excel_serial_numbers(tmp: Path):
    rows = [["sku", "fecha", "cantidad"]]
    for s, d, v in B:
        rows.append([s, (d - dt.date(1899, 12, 30)).days, int(v)])
    return write_xlsx(tmp / "serial.xlsx", {"Hoja1": rows})


@case(clean=clean_of(B), findings=["date_timezone"], fixes=["read_dates"], mapping=STD)
def date_iso_with_timezone(tmp: Path):
    return write_csv(tmp / "tz.csv", _std_rows(dfmt=lambda d: f"{d.isoformat()}T00:00:00-05:00"))


@case(clean=clean_of(B), findings=["date_two_digit_year"], fixes=["read_dates"], mapping=STD)
def date_dayfirst_two_digit_year(tmp: Path):
    return write_csv(tmp / "yy.csv", _std_rows(
        dfmt=lambda d: f"{d.day:02d}/{d.month:02d}/{d.year % 100:02d}"))


@case(asks=["date_rows_unreadable"], answers={"date_rows_unreadable": "ignore"},
      fixes=["read_dates", "ignore_unreadable_rows"], mapping=STD, clean=clean_of(B),
      extra={"unreadable": True})
def date_unreadable_rows(tmp: Path):
    rows = _std_rows()
    rows.insert(5, ["Bebidas", "n/a", ""])
    rows.insert(20, ["Snacks", "sin fecha", ""])
    return write_csv(tmp / "unreadable.csv", rows)


@case(verdict="unusable", findings=["date_missing"],
      mapping={})
def date_column_missing(tmp: Path):
    rows = [["sku", "cantidad", "precio"]] + [[s, int(v), 5] for s, _, v in B[:40]]
    return write_csv(tmp / "nodate.csv", rows)


# ── numbers ──────────────────────────────────────────────────────────────────

_FRAC = [(s, d, v + 0.5) for s, d, v in B]


@case(clean=clean_of(_FRAC), findings=["number_decimal_comma"], fixes=["read_numbers"],
      mapping=STD)
def number_decimal_comma_semicolon(tmp: Path):
    rows = _std_rows(_FRAC, vfmt=lambda v: f"{v:.1f}".replace(".", ","))
    return write_csv(tmp / "comma.csv", rows, sep=";")


_BIG = [(s, d, v * 1000 + 234.5) for s, d, v in B]


@case(clean=clean_of(_BIG), findings=["number_decimal_comma"], fixes=["read_numbers"], mapping=STD)
def number_european_thousands(tmp: Path):
    def eu(v):
        whole, frac = f"{v:,.1f}".split(".")
        return whole.replace(",", ".") + "," + frac
    return write_csv(tmp / "eu.csv", _std_rows(_BIG, vfmt=eu), sep=";")


_THOU = [(s, d, float(1000 + (int(v) * 37) % 8000)) for s, d, v in B]


@case(clean=clean_of(_THOU), asks=["number_format_ambiguous"],
      answers={"number_format_ambiguous": "."}, fixes=["read_numbers"], mapping=STD)
def number_ambiguous_thousands(tmp: Path):
    return write_csv(tmp / "thou.csv", _std_rows(_THOU, vfmt=lambda v: f"{int(v):,}"))


@case(clean=clean_of(B), asks=["demand_column_choice"],
      answers={"demand_column_choice": "importe"}, fixes=["read_numbers"],
      mapping={"sku": "sku", "date": "fecha", "demand": "importe"},
      extra={"reason": "units_or_money"})
def demand_with_currency_symbols(tmp: Path):
    rows = [["sku", "fecha", "importe"]] + [[s, d.isoformat(), f"${int(v):,}.00"]
                                              for s, d, v in B]
    return write_csv(tmp / "money.csv", rows)


@case(clean=clean_of(B), asks=["demand_column_choice"], answers={"demand_column_choice": "A"},
      mapping={"sku": "Producto", "date": "Fecha", "demand": "A"},
      extra={"reason": "units_or_money"})
def units_vs_money_generic_names(tmp: Path):
    rows = [["Producto", "Fecha", "A", "B", "C"]]
    for s, d, v in B:
        rows.append([s, d.isoformat(), int(v), 7, int(v) * 7])
    return write_csv(tmp / "generic.csv", rows)


@case(verdict="unusable", findings=["demand_constant"])
def all_zero_demand(tmp: Path):
    return write_csv(tmp / "zeros.csv", _std_rows([(s, d, 0.0) for s, d, _ in B]))


_CUM = []
for _i, _s in enumerate(SKUS):
    _run = 0
    for _k in range(26):
        _d = float(1 + (_k * 3 + _i) % 7)
        _run += _d
        _CUM.append((_s, START + dt.timedelta(days=7 * _k), _run, _d))


@case(clean=[(s, d.isoformat(), float(inc)) for s, d, _, inc in _CUM],
      asks=["demand_cumulative"], answers={"demand_cumulative": "undo"},
      fixes=["undo_cumulative"], mapping=STD)
def cumulative_running_total(tmp: Path):
    return write_csv(tmp / "cum.csv", _std_rows([(s, d, run) for s, d, run, _ in _CUM]))


# ── product codes ────────────────────────────────────────────────────────────

_ZSKUS = ["00123", "00456", "07890", "01234"]
_ZB = base(skus=_ZSKUS)


@case(clean=clean_of(_ZB), findings=["sku_leading_zeros"], fixes=["keep_leading_zeros"],
      mapping=STD)
def sku_leading_zeros_in_csv(tmp: Path):
    return write_csv(tmp / "zeros_sku.csv", _std_rows(_ZB))


_LOST = ["1234", "2345", "3456", "4567", "5678", "0678", "0789"]
_LB = base(skus=_LOST, weeks=8)
_LB_NUM = [(int(s), d, v) for s, d, v in _LB]


@case(clean=clean_of(_LB), asks=["sku_leading_zeros_lost"], answers={"sku_leading_zeros_lost": "pad"},
      fixes=["keep_leading_zeros"], mapping=STD)
def sku_leading_zeros_already_lost(tmp: Path):
    rows = [["sku", "fecha", "cantidad"]] + [[s, d, int(v)] for s, d, v in _LB_NUM]
    return write_xlsx(tmp / "lost_zeros.xlsx", {"Hoja1": rows})


def _variant_rows():
    rows = [["sku", "fecha", "cantidad"]]
    clean = []
    for s, d, v in base(skus=SKUS[:3], weeks=10):
        raw = s
        n = len(rows)
        if n % 10 == 3:
            raw = s.lower()
        elif n % 10 == 6:
            raw = f" {s} "
        rows.append([raw, d.isoformat(), int(v)])
        clean.append((s, d.isoformat(), float(v)))
    return rows, clean


@case(clean=_variant_rows()[1], asks=["sku_variants"], answers={"sku_variants": "merge"},
      findings=["sku_whitespace"], fixes=["normalize_codes"], mapping=STD)
def sku_case_and_whitespace_variants(tmp: Path):
    return write_csv(tmp / "variants.csv", _variant_rows()[0])


@case(clean=clean_of(B), asks=["sku_column_choice"], answers={"sku_column_choice": "Item"},
      mapping={"sku": "Item", "date": "fecha", "demand": "cantidad"},
      extra={"reason": "name_vs_code"})
def sku_name_and_code_both_present(tmp: Path):
    rows = [["Item", "Articulo", "fecha", "cantidad"]]
    names = {s: f"Camiseta deportiva talla {s[-1]}" for s in SKUS}
    for s, d, v in B:
        rows.append([s, names[s], d.isoformat(), int(v)])
    return write_csv(tmp / "name_code.csv", rows)


_FSKUS = ["1001", "1002", "1003", "1004"]


@case(clean=clean_of(base(skus=_FSKUS)), findings=["sku_float_codes"], fixes=["codes_as_text"],
      mapping=STD)
def sku_numeric_codes_read_as_floats(tmp: Path):
    rows = [["sku", "fecha", "cantidad"]] + [[f"{s}.0", d.isoformat(), int(v)]
                                              for s, d, v in base(skus=_FSKUS)]
    return write_csv(tmp / "float_codes.csv", rows)


# ── unusable / other kinds of file ───────────────────────────────────────────

@case(verdict="unusable", findings=["stock_snapshot"])
def stock_snapshot_file(tmp: Path):
    rows = [["sku", "descripcion", "existencia", "ubicacion"]]
    rows += [[s, f"Prod {s}", 40 + i, "A1"] for i, s in enumerate(SKUS * 6)]
    return write_csv(tmp / "stock.csv", rows)


@case(verdict="unusable", findings=["header_not_found"])
def headerless_file(tmp: Path):
    rows = [[s, d.isoformat(), int(v)] for s, d, v in B]
    return write_csv(tmp / "noheader.csv", rows)


@case(verdict="unusable", findings=["file_empty"])
def empty_file(tmp: Path):
    p = tmp / "empty.csv"
    p.write_text("", encoding="utf-8")
    return p


_SHORT = base(skus=["S-1", "S-2", "S-3", "S-4", "S-5"], weeks=20)
_SHORT = [(s, d, v) for s, d, v in _SHORT if s in ("S-1", "S-2", "S-3") or d < START + dt.timedelta(days=7 * 4)]


@case(clean=clean_of(_SHORT), zero_fixes=True, mapping=STD, extra={"short_products": 2})
def short_history_products(tmp: Path):
    return write_csv(tmp / "short.csv", _std_rows(_SHORT))


_COMBO = [(s, d, v + 0.5) for s, d, v in base(skus=["00123", "00456", "00789"])]


@case(clean=clean_of(_COMBO),
      findings=["header_row_offset", "totals_row", "date_text_months", "number_decimal_comma",
                "sku_leading_zeros"],
      fixes=["header_row", "ignore_totals_row", "read_dates", "read_numbers", "keep_leading_zeros"],
      mapping=STD)
def everything_at_once(tmp: Path):
    rows = [["Ventas semanales"], [], ["sku", "fecha", "cantidad"]]
    for s, d, v in _COMBO:
        rows.append([s, f"{d.day:02d} {ES[d.month - 1]} {d.year}", f"{v:.1f}".replace(".", ",")])
    rows.append(["Total", None, "999,5"])
    return write_csv(tmp / "combo.csv", rows, sep=";")

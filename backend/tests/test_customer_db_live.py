"""SQL data sources against REAL customer databases: PostgreSQL 16, MySQL 8 and
MariaDB 11 (see `live_customer_db.py` for where they are and what is built in
them). Nothing here touches the StockAI app database: the service functions
that normally read the `datasets` row are pointed at an in-memory row, so this
file runs with or without the app's Postgres (`pytest --noconftest` included).

Each engine is skipped when its port does not answer. Every assertion is about
what the customer's server actually did or holds: rows still there after a
refused write, connections the SERVER counts, a statement the server is no
longer running after a timeout.
"""

from __future__ import annotations

import csv
import io
import json
import threading
import time
from decimal import Decimal

import pytest

from backend.config import settings
from backend.datasources import catalog, client, probe, schema_cache
from backend.datasources import service as svc
from backend.datasources.service import build_sql_config
from backend.datasources.values import csv_cell, json_cell
from backend.errors import AppError
from backend.tests import live_customer_db as live

pytestmark = pytest.mark.offline

TARGETS = live.targets()
IDS = [t.key for t in TARGETS]


@pytest.fixture(autouse=True)
def _private_hosts_allowed(monkeypatch):
    # The test servers are on 127.0.0.1, which the SSRF policy refuses by default.
    monkeypatch.setattr(settings, "sql_sources_allow_private_hosts", True, raising=False)
    monkeypatch.setattr(settings, "sql_sources_max_concurrent_per_tenant", 4, raising=False)
    schema_cache.clear()
    yield


_ready: set[str] = set()


@pytest.fixture(params=TARGETS, ids=IDS)
def target(request):
    t = request.param
    if not live.reachable(t.host, t.port):
        pytest.skip(f"{t.key} test database not reachable on {t.host}:{t.port}")
    if t.key not in _ready:
        live.setup_target(t)
        _ready.add(t.key)
    return t


def cfg_for(t, *, user=None, password=None, database=None, port=None, host=None, **extra):
    """A stored `sql_config`, built through the same validation the API uses."""
    return build_sql_config({
        "engine": t.engine, "host": host or t.host, "port": port or t.port,
        "database": database or t.database,
        "username": user or live.RO_USER, **extra,
    }, password if password is not None else live.RO_PASSWORD)


def rows_of(cfg, sql, **kw):
    with svc._read_only_rows(cfg, sql, tenant_id="t-live", **kw) as result:
        return [tuple(r) for r in result.fetchall()], [str(c) for c in result.keys()]


# ── A fake app layer for the service functions ────────────────────────────────

class FakeApp:
    """Stands in for the `datasets` table: one SQL source row, and a record of
    every INSERT/UPDATE the service would have written."""

    def __init__(self, cfg, tmp_path, *, status="connected"):
        self.source = {"id": "ds_live", "tenant_id": "t-live", "name": "ERP", "source_type": "sql",
                       "connection_status": status, "sql_config": cfg, "saved_query": None,
                       "description": None}
        self.rows: dict[str, dict] = {"ds_live": self.source}
        self.writes: list[tuple[str, tuple]] = []
        self.tmp_path = tmp_path

    def install(self, monkeypatch):
        monkeypatch.setattr(svc, "get_source", lambda tenant_id, source_id: self.rows.get(source_id))
        monkeypatch.setattr(svc, "execute", lambda sql, params=None: self.writes.append((sql, params)))
        monkeypatch.setattr(svc, "_track_new_sales", lambda *a, **k: None)
        from backend.datasets import service as datasets_service
        monkeypatch.setattr(datasets_service, "_enforce_dataset_size", lambda *a, **k: None)
        from backend.storage import paths
        monkeypatch.setattr(paths, "dataset_dir", lambda tenant_id, ds_id: self.tmp_path / ds_id)
        return self


# ── The staged connection test ────────────────────────────────────────────────

def _stages(result):
    return {s["stage"]: s for s in result["stages"]}


class TestStagedProbe:
    def test_read_only_login_passes_every_stage(self, target):
        result = probe.run_probe(cfg_for(target), tenant_id="t-live")
        st = _stages(result)
        assert result["ok"] is True, result
        for name in ("dns", "tcp", "auth", "select", "tables"):
            assert st[name]["status"] == "ok", (name, st[name])
        assert st["privileges"]["status"] == "ok", st["privileges"]
        assert result["write_access"]["can_write"] is False
        assert result["grant_sql"] is None
        assert result["table_count"] >= 4
        assert any("stockai_big" in name for name in result["tables_sample"])
        assert result["read_only_session"] is True

    def test_a_login_that_can_write_is_flagged_with_the_grant_to_fix_it(self, target):
        result = probe.run_probe(cfg_for(target, user=target.admin_user, password=target.admin_password),
                                 tenant_id="t-live")
        st = _stages(result)
        assert result["ok"] is True
        assert st["privileges"]["status"] == "warning"
        assert st["privileges"]["code"] == "data_source_login_can_write"
        assert result["write_access"]["can_write"] is True
        joined = "\n".join(result["grant_sql"])
        assert "SELECT" in joined and "GRANT" in joined
        dumped = json.dumps(result)
        assert "password_enc" not in dumped
        if target.admin_password not in (target.admin_user, "postgres"):
            assert target.admin_password not in dumped

    def test_wrong_password_fails_at_auth_and_never_echoes_it(self, target):
        secret = "Wr0ng-Pässw0rd-€-xyz"
        result = probe.run_probe(cfg_for(target, password=secret), tenant_id="t-live")
        st = _stages(result)
        assert result["ok"] is False and result["failed_stage"] == "auth"
        assert st["auth"]["code"] == "data_source_auth_failed"
        assert st["dns"]["status"] == "ok" and st["tcp"]["status"] == "ok"
        for later in ("privileges", "select", "tables"):
            assert st[later]["status"] == "skipped"
        assert secret not in json.dumps(result)

    def test_missing_database_is_named_as_such(self, target):
        if target.engine == "mysql":
            # A MySQL login can only be told "unknown database" when it may see
            # it exists; the admin login can.
            cfg = cfg_for(target, user=target.admin_user, password=target.admin_password,
                          database="stockai_does_not_exist")
        else:
            cfg = cfg_for(target, database="stockai_does_not_exist")
        result = probe.run_probe(cfg, tenant_id="t-live")
        assert result["failed_stage"] == "auth"
        assert _stages(result)["auth"]["code"] == "data_source_database_not_found"

    def test_closed_port_fails_at_tcp(self, target):
        result = probe.run_probe(cfg_for(target, port=1), tenant_id="t-live")
        st = _stages(result)
        assert result["failed_stage"] == "tcp"
        assert st["tcp"]["code"] in ("data_source_host_unreachable", "data_source_connect_timeout")
        assert st["auth"]["status"] == "skipped"

    def test_unknown_host_fails_at_dns(self, target):
        result = probe.run_probe(cfg_for(target, host="stockai-no-such-host.invalid"), tenant_id="t-live")
        assert result["failed_stage"] == "dns"
        assert _stages(result)["dns"]["code"] == "data_source_dns_failed"

    def test_loopback_refused_unless_the_installation_allows_private_hosts(self, target, monkeypatch):
        cfg = cfg_for(target)
        monkeypatch.setattr(settings, "sql_sources_allow_private_hosts", False)
        result = probe.run_probe(cfg, tenant_id="t-live")
        assert result["failed_stage"] == "dns"
        assert _stages(result)["dns"]["code"] == "data_source_host_not_allowed"
        assert _stages(result)["tcp"]["status"] == "skipped"
        # ... and the query path is refused the same way, before connecting.
        with pytest.raises(AppError) as exc:
            rows_of(cfg, "SELECT 1")
        assert exc.value.code == "data_source_host_not_allowed"

    def test_tls_modes(self, target):
        if target.engine == "postgresql":
            # The test server runs with ssl=off: "require" must fail AT the TLS stage.
            result = probe.run_probe(cfg_for(target, ssl_mode="require"), tenant_id="t-live")
            assert result["failed_stage"] == "tls"
            assert _stages(result)["tls"]["code"] == "data_source_tls_failed"
            prefer = probe.run_probe(cfg_for(target, ssl_mode="prefer"), tenant_id="t-live")
            assert _stages(prefer)["tls"]["code"] == "data_source_tls_not_used"
        else:
            # MySQL 8 / MariaDB 11 generate a self-signed certificate.
            required = probe.run_probe(cfg_for(target, ssl_mode="require"), tenant_id="t-live")
            assert required["ok"] is True
            assert _stages(required)["tls"]["status"] == "ok"
            assert _stages(required)["tls"]["detail"]["cipher"]
            # verify-full against the system trust store cannot accept it.
            strict = probe.run_probe(cfg_for(target, ssl_mode="verify-full"), tenant_id="t-live")
            assert strict["failed_stage"] == "tls", strict
            disabled = probe.run_probe(cfg_for(target, ssl_mode="disable"), tenant_id="t-live")
            assert disabled["ok"] is True
            assert _stages(disabled)["tls"]["code"] == "data_source_tls_disabled"


# ── Types, unicode, NULL ──────────────────────────────────────────────────────

class TestValues:
    def test_every_common_type_round_trips_exactly(self, target):
        rows, cols = rows_of(cfg_for(target), "SELECT * FROM stockai_types ORDER BY id")
        first = dict(zip(cols, rows[0]))
        second = dict(zip(cols, rows[1]))

        dec = first["dec"]
        assert isinstance(dec, Decimal)
        expected_dec = ("12345678901234567890.0123456789" if target.engine == "postgresql"
                        else "1234567890123456789012345678.0123456789")
        assert csv_cell(dec) == expected_dec
        assert json_cell(dec) == expected_dec            # a string: no float rounding

        assert first["big"] == 9223372036854775807
        assert json_cell(first["big"]) == "9223372036854775807"   # beyond 2^53
        assert csv_cell(first["big"]) == 9223372036854775807
        if target.engine == "mysql":
            assert csv_cell(first["ubig"]) == 18446744073709551615

        assert csv_cell(first["d"]) == "2026-01-31"
        assert csv_cell(first["ts"]).startswith("2026-01-31T13:45:06.123456")
        if target.engine == "postgresql":
            # timestamptz keeps its instant: 13:45:06+02 is 11:45:06 UTC.
            tz = first["tstz"]
            assert tz.utcoffset() is not None
            assert tz.astimezone(__import__("datetime").timezone.utc).hour == 11
            assert csv_cell(tz).endswith(("+00:00", "+02:00")) or "+" in csv_cell(tz)

        assert csv_cell(first["b"]) == "0x00ff10"
        assert json_cell(first["b"]) == "0x00ff10"

        j = first["j"]
        parsed = j if isinstance(j, dict) else json.loads(j)
        assert parsed == {"a": [1, 2], "b": "ü"}

        assert first["txt"] == live.UNICODE_TEXT
        assert csv_cell(first["txt"]) == live.UNICODE_TEXT
        assert first["nul"] is None and csv_cell(None) is None and json_cell(None) is None
        assert all(v is None for k, v in second.items() if k != "id")

    def test_percent_and_colon_literals_are_not_parameters(self, target):
        rows, _ = rows_of(cfg_for(target),
                          "SELECT 'a :b c' AS x, 'x%y' AS y FROM stockai_big WHERE sku LIKE 'SKU-1%' "
                          "ORDER BY id LIMIT 1")
        assert rows == [("a :b c", "x%y")]

    def test_hostile_identifier_is_quoted_by_the_schema_browser(self, target):
        cfg = cfg_for(target)
        with client.open_session(cfg, tenant_id="t-live") as session:
            listing = catalog.list_tables(session.conn, target.engine)
        hostile = [tb for tb in listing["tables"] if tb["name"] == live.HOSTILE_TABLE]
        assert len(hostile) == 1, [tb["name"] for tb in listing["tables"]]
        sql = hostile[0]["select_sql"]
        assert sql, "a hostile name must still be previewable"
        rows, cols = rows_of(cfg, sql)
        assert cols == ["order", "select"]
        assert rows == [(7, live.UNICODE_TEXT)]

    def test_columns_and_row_estimates(self, target):
        cfg = cfg_for(target)
        with client.open_session(cfg, tenant_id="t-live") as session:
            listing = catalog.list_tables(session.conn, target.engine)
            big = next(tb for tb in listing["tables"] if tb["name"] == "stockai_big")
            schema = big["schema"]
            cols = catalog.list_columns(session.conn, target.engine, schema, "stockai_big")
        assert [c["name"] for c in cols["columns"]] == ["id", "sku", "qty"]
        assert cols["columns"][0]["nullable"] is False
        # An ESTIMATE from statistics: same order of magnitude, never a COUNT(*).
        assert big["row_estimate"] is not None
        assert live.BIG_ROWS / 3 <= big["row_estimate"] <= live.BIG_ROWS * 3

    def test_schema_browser_cap_reports_truncation(self, target):
        with client.open_session(cfg_for(target), tenant_id="t-live") as session:
            listing = catalog.list_tables(session.conn, target.engine, cap=2)
        assert len(listing["tables"]) == 2 and listing["truncated"] is True


# ── Read-only, twice over ─────────────────────────────────────────────────────

class TestReadOnly:
    @pytest.mark.parametrize("statement", [
        "DELETE FROM stockai_write_target",
        "SELECT 1; DELETE FROM stockai_write_target",
        "UPDATE stockai_write_target SET id = 0",
        "WITH x AS (DELETE FROM stockai_write_target RETURNING *) SELECT * FROM x",
    ])
    def test_the_guard_refuses_writes_and_the_table_is_intact(self, target, statement):
        with pytest.raises(AppError) as exc:
            rows_of(cfg_for(target, user=target.admin_user, password=target.admin_password), statement)
        assert exc.value.code.startswith("sql_")
        assert live.admin_scalar(target, "SELECT COUNT(*) FROM stockai_write_target") == 3

    def test_the_server_refuses_a_write_that_slipped_past_the_guard(self, target, monkeypatch):
        """Second layer: even with the statement check disabled, the session is
        read-only on the SERVER. Run as the ADMIN login, which could write."""
        from backend.datasources import sql_guard
        monkeypatch.setattr(sql_guard, "validate_read_only_sql", lambda sql, engine="postgresql": sql)
        cfg = cfg_for(target, user=target.admin_user, password=target.admin_password)
        with pytest.raises(Exception) as exc:
            with svc._read_only_rows(cfg, "INSERT INTO stockai_write_target VALUES (99)", tenant_id="t-live"):
                pass
        classified = svc._classified(exc.value, cfg)
        assert classified.code == "data_source_read_only_violation", classified.code
        assert live.admin_scalar(target, "SELECT COUNT(*) FROM stockai_write_target WHERE id = 99") == 0


# ── Timeouts ──────────────────────────────────────────────────────────────────

def _slow_sql(target) -> str:
    """A long read that passes the guard (no sleep functions)."""
    if target.engine == "postgresql":
        return "SELECT count(*) FROM generate_series(1, 3000000000) g /*stockai_slow*/"
    return ("SELECT /*stockai_slow*/ COUNT(*) FROM stockai_big a "
            "JOIN stockai_big b ON a.qty = b.qty AND a.sku <> b.sku")


class TestTimeouts:
    def test_statement_timeout_cancels_the_query_on_the_server(self, target):
        cfg = cfg_for(target, statement_timeout_s=2)
        started = time.monotonic()
        with pytest.raises(Exception) as exc:
            rows_of(cfg, _slow_sql(target))
        elapsed = time.monotonic() - started
        assert svc._classified(exc.value, cfg).code == "data_source_query_timeout"
        assert elapsed < 25, elapsed
        # The server stopped working on it (not just our client giving up).
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline and live.running_statements(target, live.RO_USER, "stockai_slow"):
            time.sleep(0.5)
        assert live.running_statements(target, live.RO_USER, "stockai_slow") == 0

    def test_a_sleep_past_the_guard_is_still_cut_by_the_server(self, target, monkeypatch):
        from backend.datasources import sql_guard
        monkeypatch.setattr(sql_guard, "validate_read_only_sql", lambda sql, engine="postgresql": sql)
        cfg = cfg_for(target, statement_timeout_s=2)
        sql = "SELECT pg_sleep(30)" if target.engine == "postgresql" else "SELECT SLEEP(30)"
        started = time.monotonic()
        try:
            rows, _ = rows_of(cfg, sql)
        except Exception as e:  # noqa: BLE001
            assert svc._classified(e, cfg).code == "data_source_query_timeout"
        else:
            # MySQL's max_execution_time interrupts SLEEP() by returning 1.
            assert rows == [(1,)], rows
        assert time.monotonic() - started < 20


# ── Service paths: paging, materialize, export, row cap ───────────────────────

class TestServicePaths:
    def test_paging_through_a_large_result(self, target, tmp_path, monkeypatch):
        app = FakeApp(cfg_for(target), tmp_path).install(monkeypatch)
        sql = "SELECT id, sku FROM stockai_big ORDER BY id"
        page = svc.execute_sql_query("t-live", "ds_live", sql, limit=100, offset=90_000)
        assert page["row_count"] == 100 and page["has_more"] is True
        assert page["rows"][0]["id"] == 90_001 and page["rows"][-1]["id"] == 90_100
        tail_sql = "SELECT id, sku FROM stockai_big WHERE id > 100000 ORDER BY id"
        last = svc.execute_sql_query("t-live", "ds_live", tail_sql, limit=100, offset=99_950)
        assert last["row_count"] == 50 and last["has_more"] is False
        assert last["rows"][-1]["id"] == live.BIG_ROWS
        with pytest.raises(AppError) as exc:   # past the paging ceiling: refused, not slow
            svc.execute_sql_query("t-live", "ds_live", sql, limit=10, offset=svc.MAX_QUERY_OFFSET + 1)
        assert exc.value.code == "sql_offset_out_of_range"
        exact = svc.execute_sql_query("t-live", "ds_live", "SELECT id FROM stockai_big ORDER BY id LIMIT 100",
                                      limit=100)
        # Exactly `limit` rows is NOT "truncated" (the old check said it was).
        assert exact["has_more"] is False and exact["truncated"] is False
        assert app.writes == []

    def test_materialize_writes_exact_values_and_unicode(self, target, tmp_path, monkeypatch):
        app = FakeApp(cfg_for(target), tmp_path).install(monkeypatch)
        svc.materialize_sql_source("t-live", "u1", "ds_live",
                                   sql="SELECT id, `dec`, txt, b, nul FROM stockai_types ORDER BY id"
                                   if target.engine == "mysql" else
                                   'SELECT id, "dec", txt, b, nul FROM stockai_types ORDER BY id')
        inserted = [w for w in app.writes if w[0].lstrip().startswith("INSERT INTO datasets")]
        assert len(inserted) == 1
        new_id = inserted[0][1][0]
        text = (tmp_path / new_id / "data.csv").read_text(encoding="utf-8")
        rows = list(csv.reader(io.StringIO(text)))
        assert rows[0] == ["id", "dec", "txt", "b", "nul"]
        assert rows[1][2] == live.UNICODE_TEXT and rows[1][3] == "0x00ff10" and rows[1][4] == ""
        assert "E+" not in rows[1][1] and rows[1][1].endswith(".0123456789")

    def test_row_cap_refuses_and_leaves_nothing_behind(self, target, tmp_path, monkeypatch):
        app = FakeApp(cfg_for(target), tmp_path).install(monkeypatch)
        monkeypatch.setattr(settings, "sql_materialize_max_rows", 1_000)
        started = time.monotonic()
        with pytest.raises(AppError) as exc:
            svc.materialize_sql_source("t-live", "u1", "ds_live", sql="SELECT * FROM stockai_big")
        assert exc.value.code == "sql_result_too_large"
        # Aborting does not drain the other 199,000 rows first.
        assert time.monotonic() - started < 20
        assert not [w for w in app.writes if "INSERT INTO datasets" in w[0]]
        assert list(tmp_path.iterdir()) == []

    @pytest.mark.parametrize("fmt", ["csv", "xlsx"])
    def test_export_full_result(self, target, tmp_path, monkeypatch, fmt):
        FakeApp(cfg_for(target), tmp_path).install(monkeypatch)
        out = svc.export_sql_query("t-live", "ds_live", sql="SELECT id, sku FROM stockai_big WHERE id <= 2500",
                                   fmt=fmt)
        data = b"".join(out.chunks())
        assert out.rows == 2500
        if fmt == "csv":
            lines = data.decode("utf-8-sig").splitlines()
            assert lines[0] == "id,sku" and len(lines) == 2501
        else:
            from openpyxl import load_workbook
            ws = load_workbook(io.BytesIO(data), read_only=True).active
            assert sum(1 for _ in ws.iter_rows()) == 2501


# ── Concurrency and leaks ─────────────────────────────────────────────────────

class TestResources:
    def test_per_tenant_slots_bound_concurrent_connections(self, target, monkeypatch):
        monkeypatch.setattr(settings, "sql_sources_max_concurrent_per_tenant", 2)
        monkeypatch.setattr(client, "GATE_WAIT_S", 0.3)
        cfg = cfg_for(target, statement_timeout_s=30)
        # A query that HOLDS its slot for a known time (the guard refuses sleep
        # functions, so it is bypassed here): the test measures the gate, not
        # how fast this machine scans a table.
        from backend.datasources import sql_guard
        monkeypatch.setattr(sql_guard, "validate_read_only_sql", lambda sql, engine="postgresql": sql)
        heavy = "SELECT pg_sleep(2)" if target.engine == "postgresql" else "SELECT SLEEP(2)"
        outcomes: list[str] = []
        peak = [0]
        stop = threading.Event()

        def sample():
            while not stop.is_set():
                peak[0] = max(peak[0], live.server_connections(target, live.RO_USER))
                time.sleep(0.05)

        def worker():
            try:
                rows_of(cfg, heavy)
                outcomes.append("ok")
            except AppError as e:
                outcomes.append(e.code)

        sampler = threading.Thread(target=sample)
        sampler.start()
        threads = [threading.Thread(target=worker) for _ in range(6)]
        for th in threads:
            th.start()
        for th in threads:
            th.join(120)
        stop.set()
        sampler.join()
        assert outcomes.count("data_source_busy") >= 1, outcomes
        assert outcomes.count("ok") >= 2, outcomes
        assert peak[0] <= 2, f"server saw {peak[0]} connections from one tenant"

    def test_no_connection_or_thread_leak_after_50_mixed_operations(self, target, tmp_path, monkeypatch):
        cfg = cfg_for(target)
        FakeApp(cfg, tmp_path).install(monkeypatch)
        monkeypatch.setattr(settings, "sql_materialize_max_rows", 50)
        time.sleep(1)
        before_conn = live.server_connections(target, live.RO_USER)
        before_threads = threading.active_count()
        for i in range(50):
            kind = i % 5
            if kind == 0:
                rows_of(cfg, "SELECT id FROM stockai_big WHERE id < 10")
            elif kind == 1:
                with pytest.raises(AppError):
                    svc.execute_sql_query("t-live", "ds_live", "SELECT no_such_column FROM stockai_big")
            elif kind == 2:
                with pytest.raises(AppError):   # cap exceeded mid-stream
                    svc.materialize_sql_source("t-live", "u1", "ds_live", sql="SELECT * FROM stockai_big")
            elif kind == 3:
                svc.execute_sql_query("t-live", "ds_live", "SELECT id FROM stockai_big ORDER BY id",
                                      limit=10, offset=1000)
            else:
                probe.run_probe(cfg, tenant_id="t-live")
        deadline = time.monotonic() + 15
        while time.monotonic() < deadline and live.server_connections(target, live.RO_USER) > before_conn:
            time.sleep(0.5)
        assert live.server_connections(target, live.RO_USER) == before_conn
        assert threading.active_count() <= before_threads + 4   # the bounded DNS pool at most


# ── Credentials ───────────────────────────────────────────────────────────────

class TestCredentials:
    def test_a_password_with_non_ascii_and_url_characters_connects(self, target):
        # RO_PASSWORD is "ro-pässwörd€;@/:%": every fixture connection above
        # already proves it; assert it explicitly once.
        rows, _ = rows_of(cfg_for(target), "SELECT 1")
        assert rows == [(1,)]

    def test_a_connection_string_creates_the_same_working_config(self, target):
        from urllib.parse import quote
        scheme = "postgresql" if target.engine == "postgresql" else "mysql"
        cs = (f"{scheme}://{live.RO_USER}:{quote(live.RO_PASSWORD, safe='')}@{target.host}:{target.port}/"
              f"{target.database}")
        cfg = build_sql_config({}, None, connection_string=cs)
        assert cfg["engine"] == target.engine and cfg["port"] == target.port
        rows, _ = rows_of(cfg, "SELECT 1")
        assert rows == [(1,)]
